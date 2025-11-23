#!/usr/bin/env python3
"""
Generate PDF and DOCX versions of the executive marketing plan.
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
import re

def read_markdown(file_path):
    """Read the markdown file."""
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()

def create_pdf(content, output_path):
    """Generate PDF from markdown content."""
    doc = SimpleDocTemplate(output_path, pagesize=letter,
                           rightMargin=0.75*inch, leftMargin=0.75*inch,
                           topMargin=0.75*inch, bottomMargin=0.75*inch)

    # Container for the 'Flowable' objects
    elements = []

    # Define styles
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#1a1a1a'),
        spaceAfter=20,
        alignment=TA_CENTER,
        fontName='Helvetica-Bold'
    )

    h1_style = ParagraphStyle(
        'CustomH1',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#2c3e50'),
        spaceAfter=12,
        spaceBefore=20,
        fontName='Helvetica-Bold'
    )

    h2_style = ParagraphStyle(
        'CustomH2',
        parent=styles['Heading2'],
        fontSize=14,
        textColor=colors.HexColor('#34495e'),
        spaceAfter=10,
        spaceBefore=15,
        fontName='Helvetica-Bold'
    )

    h3_style = ParagraphStyle(
        'CustomH3',
        parent=styles['Heading3'],
        fontSize=12,
        textColor=colors.HexColor('#34495e'),
        spaceAfter=8,
        spaceBefore=12,
        fontName='Helvetica-Bold'
    )

    body_style = ParagraphStyle(
        'CustomBody',
        parent=styles['BodyText'],
        fontSize=10,
        spaceAfter=8,
        alignment=TA_JUSTIFY,
        fontName='Helvetica'
    )

    bullet_style = ParagraphStyle(
        'CustomBullet',
        parent=styles['BodyText'],
        fontSize=10,
        spaceAfter=4,
        leftIndent=20,
        fontName='Helvetica'
    )

    # Parse markdown
    lines = content.split('\n')
    i = 0
    first_heading = True

    while i < len(lines):
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        # Main title (first H1)
        if line.startswith('# ') and first_heading:
            title = line[2:].strip()
            elements.append(Paragraph(title, title_style))
            elements.append(Spacer(1, 0.3*inch))
            first_heading = False

        # H2 headings (##)
        elif line.startswith('## '):
            heading = line[3:].strip()
            if heading != '---':
                elements.append(Spacer(1, 0.2*inch))
                elements.append(Paragraph(heading, h1_style))

        # H3 headings (###)
        elif line.startswith('### '):
            heading = line[4:].strip()
            elements.append(Paragraph(heading, h2_style))

        # H4 headings (####)
        elif line.startswith('#### '):
            heading = line[5:].strip()
            elements.append(Paragraph(heading, h3_style))

        # Horizontal rules
        elif line.startswith('---'):
            elements.append(Spacer(1, 0.1*inch))

        # Bold text starting lines (like **Week 1:**)
        elif line.startswith('**') and '**' in line[2:]:
            # Convert markdown bold to HTML bold for reportlab
            text = line.replace('**', '<b>', 1).replace('**', '</b>', 1)
            elements.append(Paragraph(text, body_style))

        # Bullet points
        elif line.startswith('- '):
            bullet_text = '• ' + line[2:]
            # Convert markdown bold to HTML
            bullet_text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', bullet_text)
            elements.append(Paragraph(bullet_text, bullet_style))

        # Regular paragraphs
        else:
            # Convert markdown bold to HTML
            text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', line)
            elements.append(Paragraph(text, body_style))

        i += 1

    # Build PDF
    doc.build(elements)
    print(f"PDF generated: {output_path}")

def create_docx(content, output_path):
    """Generate Word document from markdown content."""
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # Parse markdown
    lines = content.split('\n')
    i = 0
    first_heading = True

    while i < len(lines):
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        # Main title (first H1)
        if line.startswith('# ') and first_heading:
            title = line[2:].strip()
            heading = doc.add_heading(title, level=0)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            first_heading = False

        # H2 headings (##)
        elif line.startswith('## '):
            heading_text = line[3:].strip()
            if heading_text != '---':
                doc.add_heading(heading_text, level=1)

        # H3 headings (###)
        elif line.startswith('### '):
            heading_text = line[4:].strip()
            doc.add_heading(heading_text, level=2)

        # H4 headings (####)
        elif line.startswith('#### '):
            heading_text = line[5:].strip()
            doc.add_heading(heading_text, level=3)

        # Horizontal rules
        elif line.startswith('---'):
            doc.add_paragraph()

        # Bullet points
        elif line.startswith('- '):
            bullet_text = line[2:]
            # Handle bold text in bullets
            para = doc.add_paragraph(style='List Bullet')
            add_formatted_text(para, bullet_text)

        # Regular paragraphs
        else:
            para = doc.add_paragraph()
            add_formatted_text(para, line)

        i += 1

    # Save document
    doc.save(output_path)
    print(f"Word document generated: {output_path}")

def add_formatted_text(paragraph, text):
    """Add text to paragraph with bold formatting."""
    # Split by bold markers
    parts = re.split(r'(\*\*.*?\*\*)', text)

    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            # Bold text
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            # Regular text
            paragraph.add_run(part)

def main():
    # Read the executive markdown file
    content = read_markdown('/home/user/sterile_money/MARKETING_PLAN_EXECUTIVE.md')

    # Generate PDF
    create_pdf(content, '/home/user/sterile_money/SPU_LMS_Marketing_Plan.pdf')

    # Generate Word document
    create_docx(content, '/home/user/sterile_money/SPU_LMS_Marketing_Plan.docx')

    print("\n✓ Executive marketing plan generated successfully!")
    print("  - PDF: SPU_LMS_Marketing_Plan.pdf")
    print("  - Word: SPU_LMS_Marketing_Plan.docx")

if __name__ == '__main__':
    main()
